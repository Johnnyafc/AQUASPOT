# -*- coding: utf-8 -*-
"""
================================================================================
SCRIPT PARA GOOGLE COLAB / LOCAL: CARGA DEL PLAN MAESTRO A FIRESTORE
================================================================================
Este script lee el archivo Excel "PLAN MAESTRO CARACOLES HH.xlsx" y carga
automáticamente el catálogo estructurado de actividades y repuestos a la
colección 'catalogo_actividades' de Cloud Firestore.

INSTRUCCIONES PARA GOOGLE COLAB:
1. Abre Google Colab (https://colab.research.google.com).
2. Instala las dependencias ejecutando la celda:
   !pip install openpyxl firebase-admin
3. Sube a Colab:
   - Tu archivo de credenciales de Firebase: "serviceAccountKey.json"
   - El archivo Excel: "PLAN MAESTRO CARACOLES HH.xlsx"
4. Ejecuta este script.
================================================================================
"""

import os
import sys
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
import openpyxl
import firebase_admin
from firebase_admin import credentials, firestore

def inicializar_firebase(cred_path="serviceAccountKey.json"):
    """Inicializa la app de Firebase Admin."""
    if not firebase_admin._apps:
        if os.path.exists(cred_path):
            cred = credentials.Certificate(cred_path)
            firebase_admin.initialize_app(cred)
            print(f"✅ Firebase inicializado con credencial: {cred_path}")
        else:
            # Si se ejecuta con credenciales predeterminadas de entorno
            try:
                firebase_admin.initialize_app()
                print("✅ Firebase inicializado con credenciales predeterminadas.")
            except Exception as e:
                print(f"❌ Error al inicializar Firebase: {e}")
                print("Por favor sube 'serviceAccountKey.json' a la sesión de Colab.")
                raise e
    return firestore.client()

def parsear_plan_maestro(excel_path="PLAN MAESTRO CARACOLES HH.xlsx"):
    """Lee el archivo Excel y estructura las actividades."""
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"No se encontró el archivo: {excel_path}")

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb['USO POSVENTA (2)']

    # 1. Identificar las filas de inicio de cada actividad (MO...)
    mo_rows = []
    for r in range(3, ws.max_row + 1):
        c1 = ws.cell(r, 1).value
        if c1 and str(c1).strip().startswith('MO'):
            mo_rows.append((r, str(c1).strip()))

    print(f"📋 Se detectaron {len(mo_rows)} actividades principales de Mano de Obra (MO).")

    actividades = []
    for i, (r_start, code) in enumerate(mo_rows):
        r_end = mo_rows[i+1][0] - 1 if i+1 < len(mo_rows) else ws.max_row
        
        nombre = ws.cell(r_start, 2).value or ""
        hh_val = ws.cell(r_start, 4).value
        desc_trabajo = ws.cell(r_start, 9).value or ""
        
        # Repuestos e insumos internos (Cols 1-4)
        internos = []
        for r in range(r_start + 1, r_end + 1):
            icod = ws.cell(r, 1).value
            idesc = ws.cell(r, 2).value
            iuni = ws.cell(r, 3).value
            ican = ws.cell(r, 4).value
            if icod and idesc:
                try:
                    cant = float(ican) if ican is not None else 1.0
                except:
                    cant = 1.0
                internos.append({
                    "codigo": str(icod).strip(),
                    "descripcion": str(idesc).strip(),
                    "unidad": str(iuni).strip() if iuni else "UNIDAD",
                    "cantidad": cant
                })

        # Ítems comerciales y texto formal de 'Incluye' (Cols 5-9)
        comerciales = []
        incluye_text = ""
        for r in range(r_start + 1, r_end + 1):
            ccod = ws.cell(r, 5).value
            cdesc = ws.cell(r, 6).value
            cuni = ws.cell(r, 7).value
            ccan = ws.cell(r, 8).value
            cinc = ws.cell(r, 9).value
            
            if cinc and not incluye_text and str(cinc).strip().startswith("Incluye:"):
                incluye_text = str(cinc).strip()
                
            if ccod and cdesc:
                try:
                    cant = float(ccan) if ccan is not None else 1.0
                except:
                    cant = 1.0
                comerciales.append({
                    "codigo": str(ccod).strip(),
                    "descripcion": str(cdesc).strip(),
                    "unidad": str(cuni).strip() if cuni else "UNIDAD",
                    "cantidad": cant
                })

        # Procesar Horas Hombre
        hh = None
        es_variable = True
        if hh_val is not None:
            try:
                hh = float(hh_val)
                es_variable = False
            except:
                hh = None
                es_variable = True

        doc_id = f"caracol_{code.lower()}"
        actividades.append({
            "id": doc_id,
            "codigo": code,
            "nombre": str(nombre).strip(),
            "equipo": "caracol",
            "horasHombre": hh,
            "esVariable": es_variable,
            "descripcionTrabajo": str(desc_trabajo).strip(),
            "incluye": incluye_text,
            "itemsInternos": internos,
            "itemsComerciales": comerciales,
            "activo": True
        })

    return actividades

def subir_a_firestore(actividades, db):
    """Guarda cada actividad en la colección 'catalogo_actividades' de Firestore."""
    coleccion = db.collection("catalogo_actividades")
    batch = db.batch()
    batch_count = 0
    total = len(actividades)

    print(f"🚀 Iniciando carga a Firestore en 'catalogo_actividades'...")

    for act in actividades:
        doc_ref = coleccion.document(act["id"])
        batch.set(doc_ref, act)
        batch_count += 1
        print(f"  ➜ Preparado: [{act['codigo']}] {act['nombre']} (HH: {act['horasHombre'] if not act['esVariable'] else 'Variable'})")

        if batch_count >= 400: # Firestore soporta hasta 500 ops por batch
            batch.commit()
            batch = db.batch()
            batch_count = 0

    if batch_count > 0:
        batch.commit()

    print(f"\n🎉 ¡ÉXITO TOTAL! Se han sincronizado {total} actividades en Firestore.")

if __name__ == "__main__":
    # Buscar el archivo Excel en rutas comunes
    rutas_posibles = [
        "PLAN MAESTRO CARACOLES HH.xlsx",
        "assets/documents/PLAN MAESTRO CARACOLES HH.xlsx",
        "../assets/documents/PLAN MAESTRO CARACOLES HH.xlsx",
        "/content/PLAN MAESTRO CARACOLES HH.xlsx"
    ]
    
    excel_encontrado = None
    for r in rutas_posibles:
        if os.path.exists(r):
            excel_encontrado = r
            break

    if not excel_encontrado:
        print("⚠️ No se encontró 'PLAN MAESTRO CARACOLES HH.xlsx'. Por favor verifica la ruta.")
    else:
        print(f"📂 Archivo detectado: {excel_encontrado}")
        actividades = parsear_plan_maestro(excel_encontrado)
        
        # Guardar copia local en JSON por conveniencia
        json_path = os.path.join(os.path.dirname(excel_encontrado), "plan_maestro_caracoles.json")
        try:
            import json
            with open("scripts/plan_maestro_caracoles.json", "w", encoding="utf-8") as jf:
                json.dump(actividades, jf, indent=2, ensure_ascii=False)
            print("💾 Respaldo JSON generado en: scripts/plan_maestro_caracoles.json")
        except Exception as ej:
            pass
        
        # Inicializar Firebase y subir
        try:
            db = inicializar_firebase()
            subir_a_firestore(actividades, db)
        except Exception as e:
            print(f"⚠️ Detalle: {e}")
            print(f"💡 Nota: Si aún no has colocado 'serviceAccountKey.json', puedes ejecutar la parte de parsing localmente.")
            print(f"Total actividades parseadas listas para subir: {len(actividades)}")
